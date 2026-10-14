"""Ranking intelligence built from Google Search Console data (free; no rank-tracking service).

Everything here reads the permanent history in `gsc_daily` and turns it into:

  * a dashboard: totals, position groups, 3/7/28-day comparisons
  * movers: searches that gained or lost position, new and lost searches, pages up and down
  * a keyword map: each page's target keyword against what Google actually shows
  * opportunities: low click rate, positions 4-10 and 11-20, cannibalisation, unmapped searches
  * an achievement log (rank_events) and an action plan (rank_actions), each recorded once
  * a spreadsheet export laid out for Google Sheets

Positions are Search Console averages weighted by impressions, never live ranks, and are labelled that way.
Rules are fixed and published here; nothing is guessed. Causes are only stated with a confidence level.
"""
import datetime
import hashlib
import io
import json
import os
import re

from flask import current_app

from . import content, db

BUCKETS = [("1-3", 1, 3), ("4-10", 4, 10), ("11-20", 11, 20), ("21-50", 21, 50), ("51-100", 51, 100), ("100+", 101, 10 ** 6)]
THRESHOLDS = (3, 10, 20)
RECHECK_DAYS = 14
ACTIONS_DONE_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "seo", "actions_done.json")


def _d(s):
    return datetime.date.fromisoformat(s)


def latest_date():
    r = db.query("SELECT MAX(date) d FROM gsc_daily", one=True)
    return _d(r["d"]) if r and r["d"] else None


def window(days, end=None, back=0):
    """(start, end) ISO dates for a window of `days` ending `back` windows before the latest data date."""
    end = end or latest_date()
    if not end:
        return None
    e = end - datetime.timedelta(days=days * back)
    return ((e - datetime.timedelta(days=days - 1)).isoformat(), e.isoformat())


def agg(win, by="query", country=None):
    """{key: dict(clicks, impressions, ctr, position)} for a window. Position is impression-weighted."""
    if not win:
        return {}
    col = {"query": "query", "page": "page", "query_page": "query || char(9) || page"}[by]
    sql = (f"SELECT {col} k, SUM(clicks) c, SUM(impressions) i, SUM(position * impressions) pw FROM gsc_daily "
           "WHERE date BETWEEN ? AND ?" + (" AND country = ?" if country else "") + f" GROUP BY {col}")
    args = win + ((country,) if country else ())
    out = {}
    for r in db.query(sql, args):
        imp = r["i"] or 0
        out[r["k"]] = dict(clicks=int(r["c"] or 0), impressions=int(imp),
                           ctr=round((r["c"] or 0) / imp * 100, 2) if imp else 0.0,
                           position=round(r["pw"] / imp, 1) if imp else None)
    return out


def totals(win):
    if not win:
        return None
    r = db.query("SELECT SUM(clicks) c, SUM(impressions) i, SUM(position * impressions) pw, COUNT(DISTINCT query) q, "
                 "COUNT(DISTINCT page) p FROM gsc_daily WHERE date BETWEEN ? AND ?", win, one=True)
    imp = r["i"] or 0
    return dict(start=win[0], end=win[1], clicks=int(r["c"] or 0), impressions=int(imp),
                ctr=round((r["c"] or 0) / imp * 100, 2) if imp else 0.0,
                position=round(r["pw"] / imp, 1) if imp else None, queries=r["q"] or 0, pages=r["p"] or 0)


def bucket(pos):
    if pos is None:
        return None
    for name, lo, hi in BUCKETS:
        if lo - 0.5 <= pos < hi + 0.5:
            return name
    return "100+"


# ---------------------------------------------------------------- keyword map
def keyword_map(cur=None):
    """Each published page's main keyword (first in its `keywords`) against Search Console's view of it."""
    cur = cur if cur is not None else agg(window(28), "query_page")
    by_query = {}
    for k, v in cur.items():
        q, p = k.split("\t", 1)
        by_query.setdefault(q.lower(), []).append((p, v))
    out = []
    for kind in ("guides", "glossary"):
        for doc in content.list_docs(kind):
            kws = doc.keywords or []
            if not kws:
                continue
            kw = kws[0].lower()
            rows = sorted(by_query.get(kw, []), key=lambda pv: -pv[1]["impressions"])
            if not rows:
                status, pos, imp, ranking = "not seen in Google yet", None, 0, ""
            else:
                ranking, v = rows[0]
                imp = sum(r[1]["impressions"] for r in rows)
                pos = v["position"]
                status = "ranking on the target page" if ranking == doc.path else "another page ranks for it"
            out.append(dict(keyword=kws[0], target=doc.path, title=doc.title, status=status, position=pos,
                            group=bucket(pos), impressions=imp, ranking_page=ranking, secondary=", ".join(kws[1:4])))
    out.sort(key=lambda r: (r["position"] is None, r["position"] or 0))
    return out


def mapped_keywords():
    kws = set()
    for kind in ("guides", "glossary"):
        for doc in content.list_docs(kind):
            kws.update(k.lower() for k in (doc.keywords or []))
    return kws


# ---------------------------------------------------------------- analysis
def analysis():
    """Everything the Rankings page shows. Empty-safe."""
    last = latest_date()
    if not last:
        return dict(ready=False)
    w7, p7 = window(7), window(7, back=1)
    w28, p28 = window(28), window(28, back=1)
    w3, p3 = window(3), window(3, back=1)
    q28, q7, qp7 = agg(w28), agg(w7), agg(p7)
    groups = {name: 0 for name, _, _ in BUCKETS}
    for v in q28.values():
        if v["position"] is not None:
            groups[bucket(v["position"])] += 1
    movers = []
    for q, v in q7.items():
        prev = qp7.get(q)
        if prev and v["impressions"] >= 5 and prev["impressions"] >= 5 and v["position"] and prev["position"]:
            movers.append(dict(query=q, before=prev["position"], after=v["position"],
                               change=round(prev["position"] - v["position"], 1), impressions=v["impressions"]))
    movers.sort(key=lambda m: -m["change"])
    new = sorted(({**v, "query": q} for q, v in q7.items() if q not in qp7), key=lambda r: -r["impressions"])
    lost = sorted(({**v, "query": q} for q, v in qp7.items() if q not in q7 and v["impressions"] >= 5),
                  key=lambda r: -r["impressions"])
    pg7, pgp7 = agg(w7, "page"), agg(p7, "page")
    pages = []
    for p in set(pg7) | set(pgp7):
        a, b = pg7.get(p, {}), pgp7.get(p, {})
        pages.append(dict(page=p, clicks=a.get("clicks", 0), clicks_before=b.get("clicks", 0),
                          impressions=a.get("impressions", 0), impressions_before=b.get("impressions", 0),
                          change=a.get("impressions", 0) - b.get("impressions", 0)))
    pages.sort(key=lambda r: -r["change"])
    low_ctr = sorted(({**v, "query": q} for q, v in q28.items()
                      if v["impressions"] >= 100 and v["ctr"] < 1 and v["position"] and v["position"] <= 10),
                     key=lambda r: -r["impressions"])
    near = sorted(({**v, "query": q} for q, v in q28.items()
                   if v["impressions"] >= 30 and v["position"] and 3.5 <= v["position"] < 20.5),
                  key=lambda r: r["position"])
    qp28 = agg(w28, "query_page")
    cannibal = cannibalisation(qp28)
    kmap = keyword_map(qp28)
    mapped = mapped_keywords()
    unmapped = sorted(({**v, "query": q} for q, v in q28.items() if q.lower() not in mapped and v["impressions"] >= 20),
                      key=lambda r: -r["impressions"])
    countries = _by(w28, "country")[:10]
    devices = _by(w28, "device")
    since = (last - datetime.timedelta(days=6)).isoformat()
    ev7 = {k: db.query("SELECT COUNT(*) n FROM rank_events WHERE kind=? AND period >= ?", (k, since), one=True)["n"]
           for k in ("top3", "top10", "top20")}
    open_by = {r["priority"]: r["n"] for r in db.query("SELECT priority, COUNT(*) n FROM rank_actions WHERE status='open' GROUP BY priority")}
    daily_rows = [dict(r) for r in db.query("SELECT date, SUM(clicks) clicks, SUM(impressions) impressions FROM gsc_daily "
                                             "WHERE date >= ? GROUP BY date ORDER BY date", (window(90)[0],))]
    return dict(ready=True, last=last.isoformat(), entered=ev7, open_actions=open_by, chart=_chart(daily_rows), t3=totals(w3), p3=totals(p3), t7=totals(w7), p7=totals(p7),
                t28=totals(w28), p28=totals(p28), groups=groups, tracked=len(q28),
                gains=[m for m in movers if m["change"] > 0][:15], losses=sorted([m for m in movers if m["change"] < 0],
                                                                                  key=lambda m: m["change"])[:15],
                new=new[:20], lost=lost[:20], pages_up=[p for p in pages if p["change"] > 0][:10],
                pages_down=sorted([p for p in pages if p["change"] < 0], key=lambda p: p["change"])[:10],
                low_ctr=low_ctr[:20], near=near[:30], cannibal=cannibal[:20], kmap=kmap, unmapped=unmapped[:30],
                countries=countries, devices=devices,
                daily=daily_rows)


def _chart(rows, w=720, h=220):
    """Inline SVG bar chart of daily impressions with clicks marked; no scripts needed."""
    if not rows:
        return ""
    mx = max(r["impressions"] for r in rows) or 1
    bw = w / len(rows)
    bars = []
    for i, r in enumerate(rows):
        bh = (h - 70) * r["impressions"] / mx
        x = i * bw
        bars.append(f'<rect x="{x + 1:.1f}" y="{h - 34 - bh:.1f}" width="{max(bw - 2, 1):.1f}" height="{bh:.1f}" fill="#0d6a54">'
                    f'<title>{r["date"]}: {int(r["impressions"])} impressions, {int(r["clicks"])} clicks</title></rect>')
        if r["clicks"]:
            bars.append(f'<circle cx="{x + bw / 2:.1f}" cy="{h - 42 - bh:.1f}" r="5" fill="#d99a14"/>')
    return (f'<svg viewBox="0 0 {w} {h}" width="100%" role="img" aria-label="Daily impressions from {rows[0]["date"]} to '
            f'{rows[-1]["date"]}, highest {int(mx)}">' + "".join(bars) +
            f'<text x="0" y="{h - 6}" font-size="20" fill="currentColor">{rows[0]["date"]}</text>'
            f'<text x="{w}" y="{h - 6}" font-size="20" text-anchor="end" fill="currentColor">{rows[-1]["date"]}</text>'
            f'<text x="0" y="22" font-size="20" fill="currentColor">Highest day: {int(mx)} impressions</text></svg>')


def _by(win, col):
    rows = db.query(f"SELECT {col} k, SUM(clicks) c, SUM(impressions) i, SUM(position*impressions) pw FROM gsc_daily "
                    f"WHERE date BETWEEN ? AND ? GROUP BY {col} ORDER BY i DESC", win)
    return [dict(key=r["k"], clicks=int(r["c"] or 0), impressions=int(r["i"] or 0),
                 position=round(r["pw"] / r["i"], 1) if r["i"] else None) for r in rows]


def cannibalisation(qp):
    by_q = {}
    for k, v in qp.items():
        q, p = k.split("\t", 1)
        by_q.setdefault(q, []).append((p, v))
    out = []
    for q, rows in by_q.items():
        total = sum(v["impressions"] for _, v in rows)
        strong = [(p, v) for p, v in rows if v["impressions"] >= 5 and total and v["impressions"] / total >= 0.1]
        if len(strong) >= 2:
            strong.sort(key=lambda pv: -pv[1]["impressions"])
            out.append(dict(query=q, impressions=total,
                            pages=[dict(page=p, impressions=v["impressions"], position=v["position"]) for p, v in strong]))
    out.sort(key=lambda r: -r["impressions"])
    return out


# ---------------------------------------------------------------- achievements (recorded once each)
def record_events(today=None):
    last = latest_date()
    if not last:
        return 0
    found = (today or datetime.date.today()).isoformat()
    w7, p7 = window(7), window(7, back=1)
    cur, prev = agg(w7), agg(p7)
    rows = []
    for q, v in cur.items():
        if v["impressions"] < 3 or v["position"] is None:
            continue
        before = prev.get(q, {}).get("position")
        for t in THRESHOLDS:
            if v["position"] <= t + 0.49 and (before is None or before > t + 0.49):
                rows.append((found, f"top{t}", q, "", f"{before} (avg, {p7[0]} to {p7[1]})" if before else "not shown",
                             f"{v['position']} (avg, {w7[0]} to {w7[1]})", f"{w7[0]} to {w7[1]}"))
    # First impressions and first clicks per page, from the whole history.
    for r in db.query("SELECT page, MIN(date) d FROM gsc_daily WHERE impressions > 0 GROUP BY page"):
        rows.append((found, "first_impressions", "", r["page"], "none", f"first shown on {r['d']}", r["d"]))
    for r in db.query("SELECT page, MIN(date) d FROM gsc_daily WHERE clicks > 0 GROUP BY page"):
        rows.append((found, "first_click", "", r["page"], "none", f"first click on {r['d']}", r["d"]))
    conn = db.get_db()
    cur_ = conn.executemany("INSERT OR IGNORE INTO rank_events(found_on, kind, query, page, before, after, period) "
                            "VALUES (?,?,?,?,?,?,?)", rows)
    conn.commit()
    return cur_.rowcount


EVENT_LABEL = {"top3": "Entered the top 3", "top10": "Entered the top 10", "top20": "Entered the top 20",
               "first_impressions": "Page shown in Google for the first time", "first_click": "Page received its first click"}


def events(limit=200):
    return [dict(r, label=EVENT_LABEL.get(r["kind"], r["kind"])) for r in
            db.query("SELECT * FROM rank_events ORDER BY found_on DESC, id DESC LIMIT ?", (limit,))]


# ---------------------------------------------------------------- diagnosis log and action plan (fixed rules)
def _aid(kind, query, page):
    return hashlib.sha1(f"{kind}|{query}|{page}".encode()).hexdigest()[:12]


def generate_actions(today=None):
    a = analysis()
    if not a.get("ready"):
        return 0
    found = (today or datetime.date.today())
    recheck = (found + datetime.timedelta(days=RECHECK_DAYS)).isoformat()
    period = f"{a['t28']['start']} to {a['t28']['end']}"
    qp = agg(window(28), "query_page")
    top_page = {}
    for k, v in qp.items():
        q, p = k.split("\t", 1)
        if q not in top_page or v["impressions"] > top_page[q][1]:
            top_page[q] = (p, v["impressions"])
    rows = []

    def add(kind, query, page, evidence, explanation, confidence, priority, fix, metric, baseline):
        rows.append((_aid(kind, query, page), found.isoformat(), kind, query, page, evidence, explanation, confidence,
                     priority, fix, metric, baseline, recheck))

    for r in a["low_ctr"]:
        pg = top_page.get(r["query"], ("", 0))[0]
        add("low_ctr", r["query"], pg, f"{r['impressions']} impressions, {r['clicks']} clicks ({r['ctr']}% CTR) at average position {r['position']}, {period}, Search Console.",
            "People see the result but rarely choose it; the title and description may not match what they are looking for.",
            "medium", "high", f"Rewrite the meta title and description of {pg} to answer \"{r['query']}\" directly.", "CTR for this search", f"{r['ctr']}%")
    for r in a["near"]:
        pg = top_page.get(r["query"], ("", 0))[0]
        if r["position"] <= 10.5:
            add("striking", r["query"], pg, f"Average position {r['position']} with {r['impressions']} impressions, {period}, Search Console.",
                "Already on page one; a stronger answer could reach the top three.", "medium", "high",
                f"On {pg}, add a section that answers \"{r['query']}\" in the first lines and link to it from two related pages.",
                "Average position for this search", str(r["position"]))
        else:
            add("page_two", r["query"], pg, f"Average position {r['position']} with {r['impressions']} impressions, {period}, Search Console.",
                "On page two; the page covers the topic but not fully enough to rank on page one.", "medium", "medium",
                f"Expand {pg} with the sub-topics behind \"{r['query']}\" and add internal links from related guides.",
                "Average position for this search", str(r["position"]))
    for c in a["cannibal"]:
        pages = ", ".join(f"{p['page']} ({p['impressions']} impressions, position {p['position']})" for p in c["pages"])
        add("cannibal", c["query"], c["pages"][0]["page"], f"Two or more of our pages are shown for this search: {pages}, {period}, Search Console.",
            "Google is splitting this search between our pages, which can hold both back.", "high", "high",
            f"Make {c['pages'][0]['page']} the main page for \"{c['query']}\", link to it from the others with that phrase, and focus the others on their own keywords.",
            "Share of impressions on the main page", f"{c['pages'][0]['impressions']} of {c['impressions']}")
    for m in a["losses"]:
        if m["change"] <= -3:
            pg = top_page.get(m["query"], ("", 0))[0]
            add("decline", m["query"], pg, f"Average position moved from {m['before']} to {m['after']} week on week ({m['impressions']} impressions), Search Console.",
                "The cause is not known from this data; it may be normal movement for a new site.", "low", "medium",
                f"Check whether {pg} changed recently and compare it with what now ranks for \"{m['query']}\"; recheck before making changes.",
                "Average position for this search", str(m["after"]))
    for k in a["kmap"]:
        if k["status"] == "another page ranks for it" and k["impressions"] >= 20:
            add("mapping", k["keyword"], k["target"], f"{k['ranking_page']} is shown for \"{k['keyword']}\" instead of {k['target']} ({k['impressions']} impressions), Search Console.",
                "The page written for this keyword is not the one Google prefers.", "medium", "medium",
                f"Link from {k['ranking_page']} to {k['target']} using the phrase \"{k['keyword']}\", or move the keyword to the page Google prefers.",
                "Page shown for this keyword", k["ranking_page"])
    for u in a["unmapped"]:
        if u["impressions"] >= 50 and (u["position"] or 100) > 20:
            add("new_page", u["query"], "", f"{u['impressions']} impressions at average position {u['position']}, no page targets it, {period}, Search Console.",
                "There is search demand our pages only touch on.", "medium", "medium",
                f"Add \"{u['query']}\" to the content calendar as a new page, after the duplicate check.", "Impressions for this search", str(u["impressions"]))
    conn = db.get_db()
    cur = conn.executemany("INSERT OR IGNORE INTO rank_actions(id, found_on, kind, query, page, evidence, explanation, confidence, "
                           "priority, fix, metric, baseline, recheck_on) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    conn.commit()
    return cur.rowcount


def apply_done_file():
    """Actions the automatic runs say they carried out (seo/actions_done.json: [{"id", "date", "note"}])."""
    try:
        with open(ACTIONS_DONE_FILE, encoding="utf-8") as fh:
            items = json.load(fh)
    except (OSError, ValueError):
        return 0
    n = 0
    for it in items if isinstance(items, list) else []:
        if isinstance(it, dict) and re.match(r"^[0-9a-f]{12}$", str(it.get("id", ""))):
            n += db.get_db().execute("UPDATE rank_actions SET status='done', status_note=? WHERE id=? AND status='open'",
                                     (f"Done by Claude on {it.get('date', '')}: {it.get('note', '')}"[:300], it["id"])).rowcount
    db.get_db().commit()
    return n


def follow_up(today=None):
    """After the recheck date, record what the metric shows now."""
    today = (today or datetime.date.today()).isoformat()
    w = window(28)
    if not w:
        return 0
    q = agg(w)
    n = 0
    for r in db.query("SELECT * FROM rank_actions WHERE recheck_on <= ? AND result IS NULL AND status != 'dismissed'", (today,)):
        v = q.get(r["query"])
        if r["kind"] == "low_ctr":
            now = f"{v['ctr']}%" if v else "no data"
        elif r["kind"] == "new_page":
            now = str(v["impressions"]) if v else "no data"
        elif r["kind"] in ("cannibal", "mapping"):
            continue
        else:
            now = str(v["position"]) if v else "not shown"
        db.execute("UPDATE rank_actions SET result=? WHERE id=?", (f"{r['baseline']} then, {now} now ({w[0]} to {w[1]})", r["id"]))
        n += 1
    return n


def actions(status=None):
    order = "CASE priority WHEN 'critical' THEN 0 WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END, found_on DESC"
    if status:
        return db.query(f"SELECT * FROM rank_actions WHERE status=? ORDER BY {order}", (status,))
    return db.query(f"SELECT * FROM rank_actions ORDER BY status='open' DESC, {order}")


def set_status(aid, status, note=""):
    if status not in ("open", "done", "dismissed"):
        return False
    return db.get_db().execute("UPDATE rank_actions SET status=?, status_note=? WHERE id=?", (status, note[:300], aid)).rowcount > 0


# ---------------------------------------------------------------- daily job and the snapshot the automatic runs read
def daily_job():
    events_n = record_events()
    actions_n = generate_actions()
    done_n = apply_done_file()
    follow_n = follow_up()
    snap = snapshot()
    current_app.logger.info("rankings-snapshot %s", json.dumps(snap, separators=(",", ":"))[:9000])
    return dict(events=events_n, actions=actions_n, done=done_n, followed=follow_n)


def snapshot():
    a = analysis()
    if not a.get("ready"):
        return dict(ready=False)
    return dict(ready=True, last=a["last"], t7=a["t7"], p7=a["p7"], t28=a["t28"], groups=a["groups"],
                gains=a["gains"][:5], losses=a["losses"][:5], new=[n["query"] for n in a["new"][:10]],
                actions=[dict(id=r["id"], kind=r["kind"], query=r["query"], page=r["page"], priority=r["priority"],
                              fix=r["fix"]) for r in actions("open")[:15]],
                achievements=[dict(kind=e["kind"], query=e["query"], page=e["page"], after=e["after"]) for e in events(10)])


# ---------------------------------------------------------------- spreadsheet export (Google Sheets ready)
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    a = analysis()
    wb = Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="D8EBE4")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    note = "Positions are Search Console averages weighted by impressions, not live Google ranks."
    if a.get("ready"):
        dash = [("Data up to", a["last"]), ("Searches tracked (28 days)", a["tracked"]), ("Note", note)]
        for label, t in (("Last 28 days", a["t28"]), ("Previous 28 days", a["p28"]), ("Last 7 days", a["t7"]),
                         ("Previous 7 days", a["p7"]), ("Last 3 days", a["t3"]), ("Previous 3 days", a["p3"])):
            if t:
                dash.append((f"{label} ({t['start']} to {t['end']})",
                             f"{t['clicks']} clicks, {t['impressions']} impressions, CTR {t['ctr']}%, avg position {t['position']}"))
        for g, n in a["groups"].items():
            dash.append((f"Searches at average position {g}", n))
        sheet("SEO Dashboard", ["Measure", "Value"], dash, first=True)
        sheet("Keyword Map", ["Keyword", "Target page", "Status", "Avg position (28d)", "Group", "Impressions (28d)",
                              "Page Google shows", "Secondary keywords"],
              [(k["keyword"], k["target"], k["status"], k["position"], k["group"], k["impressions"], k["ranking_page"],
                k["secondary"]) for k in a["kmap"]])
        sheet("Movers", ["Search", "Avg position before (7d)", "Avg position now (7d)", "Change", "Impressions (7d)"],
              [(m["query"], m["before"], m["after"], m["change"], m["impressions"]) for m in a["gains"] + a["losses"]])
        sheet("Opportunities", ["Type", "Search", "Impressions (28d)", "Clicks", "CTR %", "Avg position"],
              [("Low CTR", r["query"], r["impressions"], r["clicks"], r["ctr"], r["position"]) for r in a["low_ctr"]] +
              [("Positions 4-20", r["query"], r["impressions"], r["clicks"], r["ctr"], r["position"]) for r in a["near"]] +
              [("No page targets it", r["query"], r["impressions"], r["clicks"], r["ctr"], r["position"]) for r in a["unmapped"]])
        sheet("Cannibalisation", ["Search", "Total impressions (28d)", "Pages shown (impressions, avg position)"],
              [(c["query"], c["impressions"], "; ".join(f"{p['page']} ({p['impressions']}, {p['position']})" for p in c["pages"]))
               for c in a["cannibal"]])
    else:
        sheet("SEO Dashboard", ["Measure", "Value"], [("Status", "No Search Console data yet"), ("Note", note)], first=True)
    sheet("Action Plan", ["ID", "Found", "Type", "Search", "Page", "Evidence", "Likely explanation", "Confidence", "Priority",
                          "Recommended fix", "Metric", "Baseline", "Recheck on", "Status", "Status note", "Follow-up result"],
          [(r["id"], r["found_on"], r["kind"], r["query"], r["page"], r["evidence"], r["explanation"], r["confidence"],
            r["priority"], r["fix"], r["metric"], r["baseline"], r["recheck_on"], r["status"], r["status_note"], r["result"])
           for r in actions()])
    sheet("Achievements", ["Found", "Achievement", "Search", "Page", "Before", "After", "Period"],
          [(e["found_on"], e["label"], e["query"], e["page"], e["before"], e["after"], e["period"]) for e in events(1000)])
    sheet("GSC Performance", ["Date", "Search", "Page", "Country", "Device", "Clicks", "Impressions", "CTR", "Avg position"],
          [(r["date"], r["query"], r["page"], r["country"], r["device"], r["clicks"], r["impressions"], round(r["ctr"] * 100, 2),
            round(r["position"], 1)) for r in db.query("SELECT * FROM gsc_daily ORDER BY date DESC, impressions DESC LIMIT 50000")])
    from . import seoauto
    log = []
    for d in seoauto.run_dates():
        r = seoauto.run(d)
        if r:
            log.append((d, "Daily SEO run", r["summary"]))
    for w in seoauto.watch_list(200):
        log.append((w["date"], "Morning self-check (" + w["status"] + ")", w["summary"]))
    from . import gsc
    st = gsc.status()
    if st:
        log.append((datetime.datetime.utcfromtimestamp(st.get("checked_at") or 0).date().isoformat(),
                    "Search Console fetch", "ok" if st.get("ok") else "failed: " + (st.get("error") or "")))
    log.sort(reverse=True)
    sheet("Change Log", ["Date", "Event", "Detail"], log)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
